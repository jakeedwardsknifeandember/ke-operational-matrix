import streamlit as st
import datetime
import pandas as pd
import google.generativeai as genai
import gspread
import os
from fpdf import FPDF
import json
import docx
from streamlit_drawable_canvas import st_canvas
from PIL import Image
import numpy as np
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
import tempfile

# --- FIXED SYSTEM CONSTANTS ---
DRAFT_FILE_PATH = "audit_draft_backup.json"
PRIMARY_RECIPIENT_EMAIL = "jakeedwards.knifeandember@gmail.com"

# --- SAFE SECRETS HELPER ---
def get_secret(key, default=None):
    try:
        return st.secrets[key]
    except Exception:
        return default

# --- TEXT SANITIZER FOR PDF (PREVENTS '?' ENCODING GLITCHES) ---
def clean_unicode_text(text):
    if not text:
        return ""
    replacements = {
        '\u2014': '-',  # em-dash
        '\u2013': '-',  # en-dash
        '\u201c': '"',  # left double quote
        '\u201d': '"',  # right double quote
        '\u2018': "'",  # left single quote
        '\u2019': "'",  # right single quote
        '\u2022': '*',  # bullet
        '\u2026': '...',# ellipsis
    }
    for original, replacement in replacements.items():
        text = text.replace(original, replacement)
    return text.encode('latin-1', 'replace').decode('latin-1')

# --- SECURE CLOUD & AI SETUP ---
gemini_key = get_secret("GEMINI_API_KEY")
if gemini_key:
    genai.configure(api_key=gemini_key)
    model = genai.GenerativeModel('gemini-2.5-flash')
else:
    model = None

gc = None
gcp_secret = get_secret("gcp_service_account")
if gcp_secret:
    try:
        credentials = json.loads(gcp_secret)
        gc = gspread.service_account_from_dict(credentials)
    except Exception:
        gc = None
    
# --- THE SETTINGS ---
BASE_SCORE = 1000

# --- DYNAMIC CHECKLIST LOADER (FROM JSON FILES) ---
def load_all_checklists():
    checklists = {}
    checklists_dir = "checklists"
    
    if os.path.exists(checklists_dir) and os.path.isdir(checklists_dir):
        for filename in os.listdir(checklists_dir):
            if filename.endswith(".json"):
                file_path = os.path.join(checklists_dir, filename)
                try:
                    with open(file_path, "r", encoding="utf-8") as f:
                        data = json.load(f)
                        concept_name = data.get("concept_name", filename)
                        checklists[concept_name] = data
                except Exception as e:
                    st.error(f"Error loading checklist '{filename}': {e}")
                    
    if not checklists:
        checklists["Default Checklist"] = {
            "concept_name": "Default",
            "folder_slug": "default",
            "sheet_name": "Audit_Database",
            "checklist": {
                "Module 1: General Hygiene": [
                    {"Fail?": False, "ID": "1.1", "Description": "General sanitation observed.", "Class": "L1", "Notes": ""}
                ]
            }
        }
    return checklists

CHECKLIST_LIBRARY = load_all_checklists()

# --- APP MEMORY (THE MEMORY SHIELD) ---
if 'custom_findings' not in st.session_state:
    st.session_state.custom_findings = [{"note": "", "level": "None (No deduction)", "changelog": False}]
if 'photo_evidence' not in st.session_state:
    st.session_state.photo_evidence = []
if 'restored_module_states' not in st.session_state:
    st.session_state.restored_module_states = {}
if 'logged_in' not in st.session_state:
    st.session_state.logged_in = False
if 'report_generated' not in st.session_state:
    st.session_state.report_generated = False
if 'cached_report_text' not in st.session_state:
    st.session_state.cached_report_text = ""
if 'cached_score_data' not in st.session_state:
    st.session_state.cached_score_data = {}

def add_custom_finding():
    st.session_state.custom_findings.append({"note": "", "level": "None (No deduction)", "changelog": False})

def add_photo_slot():
    st.session_state.photo_evidence.append({"caption": "", "file": None})

# --- DRAFT BACKUP HELPERS ---
def save_audit_draft(concept_name, est_name, branch_loc, fsco, date_val, notes, custom_f, edited_mods):
    try:
        module_snapshots = {}
        for mod_name, mod_df in edited_mods.items():
            if isinstance(mod_df, pd.DataFrame):
                module_snapshots[mod_name] = mod_df.to_dict(orient="records")

        draft_payload = {
            "timestamp": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "concept_name": concept_name,
            "establishment_name": est_name,
            "branch_name": branch_loc,
            "fsco_name": fsco,
            "audit_date": str(date_val),
            "auditor_notes": notes,
            "custom_findings": custom_f,
            "module_states": module_snapshots
        }
        
        with open(DRAFT_FILE_PATH, "w", encoding="utf-8") as f:
            json.dump(draft_payload, f, indent=2)
        return True
    except Exception as e:
        st.error(f"Could not save draft backup: {e}")
        return False

def clear_audit_draft():
    if os.path.exists(DRAFT_FILE_PATH):
        try:
            os.remove(DRAFT_FILE_PATH)
        except Exception:
            pass

# --- TARGETED SOP READER (RAG FILTERED BY FAILED VIOLATIONS) ---
def read_company_standards(concept_folder, failed_items=None):
    target_path = os.path.join("standards", concept_folder)
    
    if os.path.exists(target_path) and os.path.isdir(target_path):
        scan_folder = target_path
    elif os.path.exists("standards") and os.path.isdir("standards"):
        scan_folder = "standards"
    else:
        return "WARNING: No standards directory found for this establishment."

    try:
        found_docs = False
        all_paragraphs = []

        for filename in os.listdir(scan_folder):
            if filename.endswith(".docx"):
                found_docs = True
                doc = docx.Document(os.path.join(scan_folder, filename))
                for p in doc.paragraphs:
                    text = p.text.strip()
                    if text:
                        all_paragraphs.append((filename, text))
        
        if not found_docs:
            return f"WARNING: No .docx SOP files found in '{scan_folder}'."

        if not failed_items:
            full_standards = ""
            current_file = ""
            for filename, text in all_paragraphs:
                if filename != current_file:
                    current_file = filename
                    full_standards += f"\n--- {filename} ---\n"
                full_standards += text + "\n"
            return full_standards

        keywords = set()
        stopwords = {"the", "and", "for", "with", "that", "this", "from", "have", "were", "been", "none", "notes", "custom", "finding", "l1", "l2", "l3"}
        for item in failed_items:
            words = [w.lower().strip(".,()[]:;\"'-") for w in item.split()]
            meaningful = [w for w in words if len(w) > 2 and w not in stopwords]
            keywords.update(meaningful)

        relevant_blocks = []
        for filename, text in all_paragraphs:
            text_lower = text.lower()
            if any(kw in text_lower for kw in keywords):
                relevant_blocks.append(f"[{filename}] {text}")

        if not relevant_blocks:
            full_standards = ""
            current_file = ""
            for filename, text in all_paragraphs:
                if filename != current_file:
                    current_file = filename
                    full_standards += f"\n--- {filename} ---\n"
                full_standards += text + "\n"
            return full_standards

        return "\n".join(relevant_blocks)
            
    except Exception as e:
        return f"WARNING: Could not read standards folder. Error: {e}"

# --- AUTOMATED EMAIL BACKEND DISPATCH ---
def auto_email_report(recipient_email, pdf_path, client_name, branch_name, score, status):
    smtp_user = get_secret("smtp_username")
    smtp_server_val = get_secret("smtp_server")
    smtp_port_val = get_secret("smtp_port")
    smtp_pass = get_secret("smtp_password")

    if not all([smtp_user, smtp_server_val, smtp_port_val, smtp_pass]):
        st.error("Missing SMTP credentials in secrets configuration.")
        return False

    try:
        msg = MIMEMultipart()
        msg['From'] = f"Knife & Ember <{smtp_user}>"
        msg['To'] = recipient_email
        full_title = f"{client_name} - {branch_name}" if branch_name else client_name
        msg['Subject'] = f"Food Safety Audit Report: {full_title} - {score:.2f}% ({status})"
        
        body = f"""Dear Management,

Please find attached the official FSCO Monthly Surveillance & Verification Report for {full_title}.

Audit Execution Date: {datetime.date.today()}
Final Verification Score: {score:.2f}%
Current Operational Status: {status}

This document serves as an official record of operational safety metrics. Required corrective actions (CAPA) must be initiated within the mandated 24-48 hour parameters.

Best regards,
Jake-Edwards L. Yboa, FSCO
Lead Auditor | Knife & Ember Food Consultancy Services
"""
        msg.attach(MIMEText(body, 'plain'))
        
        with open(pdf_path, "rb") as attachment:
            part = MIMEBase('application', 'octet-stream')
            part.set_payload(attachment.read())
            encoders.encode_base64(part)
            part.add_header('Content-Disposition', f"attachment; filename={os.path.basename(pdf_path)}")
            msg.attach(part)
            
        server = smtplib.SMTP(smtp_server_val, int(smtp_port_val))
        server.starttls()
        server.login(smtp_user, smtp_pass)
        server.send_message(msg)
        server.quit()
        return True
    except Exception as e:
        st.error(f"Failed to deliver automated email: {e}")
        return False

# ==========================================
# THE GATEKEEPER (LOGIN SCREEN)
# ==========================================
if not st.session_state.logged_in:
    st.title("Dynamic FSMS System")
    st.subheader("Access Restricted")
    st.write("Please enter the system password to access the audit tools and dashboard.")
    
    password_input = st.text_input("Password", type="password")
    
    if st.button("Login"):
        app_pass = get_secret("APP_PASSWORD")
        if app_pass is None:
            st.error("`APP_PASSWORD` missing from Streamlit secrets. Please set it in secrets.toml or Streamlit Cloud.")
        elif password_input == app_pass:
            st.session_state.logged_in = True
            st.rerun()
        else:
            st.error("Incorrect password. Please try again.")

# ==========================================
# THE MAIN APP (HIDDEN BEHIND LOGIN)
# ==========================================
if st.session_state.logged_in:
    
    st.sidebar.markdown("### Knife & Ember")
    st.sidebar.markdown("*Food Consultancy Services*")
    st.sidebar.markdown("---")
    
    st.sidebar.subheader("System Control Panel")
    
    # Store Module Selector
    concept_keys = list(CHECKLIST_LIBRARY.keys())
    default_concept_idx = 0
    if 'restored_concept' in st.session_state and st.session_state.restored_concept in concept_keys:
        default_concept_idx = concept_keys.index(st.session_state.restored_concept)

    selected_concept_name = st.sidebar.selectbox(
        "Select Checklist Profile / Module:",
        options=concept_keys,
        index=default_concept_idx
    )
    
    selected_profile = CHECKLIST_LIBRARY[selected_concept_name]
    active_folder_slug = selected_profile.get("folder_slug", "default")
    active_sheet_name = selected_profile.get("sheet_name", "Audit_Database")
    active_checklist = selected_profile.get("checklist", {})
    
    est_default = st.session_state.get('restored_est_name', "Hyeongje Grill")
    branch_default = st.session_state.get('restored_branch_name', "SM Megamall")
    fsco_default = st.session_state.get('restored_fsco_name', "Jake-Edwards L. Yboa")
    
    establishment_name = st.sidebar.text_input("Brand / Establishment Name:", value=est_default)
    branch_name = st.sidebar.text_input("Branch / Location:", value=branch_default)
    fsco_name = st.sidebar.text_input("Lead Auditor / FSCO:", value=fsco_default)
    st.sidebar.caption("Certified Food Safety Compliance Officer")
    audit_date = st.sidebar.date_input("Audit Operational Date:", datetime.date.today())
    
    st.sidebar.markdown("---")
    st.sidebar.subheader("Cloud Dispatch Routing")
    st.sidebar.info(f"Report Destination Email:\n`{PRIMARY_RECIPIENT_EMAIL}`")
    
    st.sidebar.markdown("---")
    st.sidebar.subheader("Offline Session Recovery")
    
    st.title("Operational Surveillance Suite")
    st.markdown(f"**Active Client:** {establishment_name} - {branch_name}")
    st.caption(f"Active FSMS Standards Directory: standards/{active_folder_slug}/ | Target Database: {active_sheet_name}")
    st.divider()
    
    tab1, tab2 = st.tabs(["Conduct Operational Audit", "Portfolio Analytics Dashboard"])

    with tab1:
        # --- UNSAVED DRAFT RECOVERY BANNER ---
        if os.path.exists(DRAFT_FILE_PATH):
            try:
                with open(DRAFT_FILE_PATH, "r", encoding="utf-8") as f:
                    draft_data = json.load(f)
                d_time = draft_data.get("timestamp", "Unknown time")
                d_est = draft_data.get("establishment_name", "Unsaved Establishment")
                d_br = draft_data.get("branch_name", "")
                d_label = f"{d_est} ({d_br})" if d_br else d_est
                
                st.warning(f"Unsaved audit draft detected for **{d_label}** (Saved on: {d_time}).")
                d_col1, d_col2 = st.columns([1, 4])
                with d_col1:
                    if st.button("Restore Draft", use_container_width=True):
                        st.session_state.restored_concept = draft_data.get("concept_name", selected_concept_name)
                        st.session_state.restored_est_name = draft_data.get("establishment_name", establishment_name)
                        st.session_state.restored_branch_name = draft_data.get("branch_name", branch_name)
                        st.session_state.restored_fsco_name = draft_data.get("fsco_name", fsco_name)
                        st.session_state.custom_findings = draft_data.get("custom_findings", [])
                        st.session_state.restored_notes = draft_data.get("auditor_notes", "")
                        st.session_state.restored_module_states = draft_data.get("module_states", {})
                        st.success("Draft restored successfully!")
                        st.rerun()
                with d_col2:
                    if st.button("Discard Draft", use_container_width=True):
                        clear_audit_draft()
                        st.toast("Saved draft backup cleared.")
                        st.rerun()
            except Exception as e:
                st.error(f"Error reading draft backup: {e}")

        st.subheader("Operational Checkpoints")
        st.write("Toggle deviations observed across processing corridors below:")

        edited_modules = {}

        for module_name, checkpoints in active_checklist.items():
            with st.expander(f"{module_name}", expanded=False):
                if module_name in st.session_state.restored_module_states:
                    df = pd.DataFrame(st.session_state.restored_module_states[module_name])
                else:
                    df = pd.DataFrame(checkpoints)

                edited_df = st.data_editor(
                    df,
                    column_config={
                        "Fail?": st.column_config.CheckboxColumn("Fail?", default=False),
                        "ID": st.column_config.TextColumn("Ref ID", disabled=True),
                        "Description": st.column_config.TextColumn("Checkpoint", disabled=True),
                        "Class": st.column_config.TextColumn("Class", disabled=True),
                        "Notes": st.column_config.TextColumn("Notes / Evidence")
                    },
                    hide_index=True,
                    use_container_width=True,
                    key=f"{selected_concept_name}_{module_name}"
                )
                edited_modules[module_name] = edited_df

        # Quick Save Button in Sidebar
        if st.sidebar.button("Save Draft Progress", use_container_width=True):
            current_notes = st.session_state.get("auditor_notes_field", "")
            if save_audit_draft(selected_concept_name, establishment_name, branch_name, fsco_name, audit_date, current_notes, st.session_state.custom_findings, edited_modules):
                st.sidebar.success("Draft progress saved locally!")

        st.divider()

        st.subheader("Add Custom Findings (On-the-Fly)")
        for i, finding in enumerate(st.session_state.custom_findings):
            col1, col2, col3 = st.columns([3, 2, 1])
            with col1:
                finding["note"] = st.text_input(f"Violation Note #{i+1}", value=finding["note"], key=f"note_{i}")
            with col2:
                options = ["None (No deduction)", "L1 Critical (-25 pts)", "L2 Major (-10 pts)", "L3 Minor (-2 pts)"]
                finding["level"] = st.selectbox("Risk Level", options, index=options.index(finding["level"]), key=f"lvl_{i}")
            with col3:
                st.write(" ")
                st.write(" ")
                finding["changelog"] = st.checkbox("Add to Changelog", value=finding["changelog"], key=f"log_{i}")

        st.button("Add Another Custom Finding", on_click=add_custom_finding)
        st.divider()

        # --- PHOTO EVIDENCE ATTACHMENT VAULT ---
        st.subheader("On-Site Photo Evidence Vault")
        st.write("Upload or capture photos of observations to embed directly into the report:")

        if not st.session_state.photo_evidence:
            st.session_state.photo_evidence = [{"caption": "", "file": None}]

        uploaded_photos_data = []
        for p_idx, photo_item in enumerate(st.session_state.photo_evidence):
            st.markdown(f"**Photo Evidence #{p_idx+1}**")
            p_col1, p_col2 = st.columns([3, 2])
            
            with p_col1:
                input_mode = st.radio(
                    f"Capture Method #{p_idx+1}:",
                    ["File Upload", "Camera Snap"],
                    index=0,
                    horizontal=True,
                    key=f"photo_src_{p_idx}"
                )
                if input_mode == "Camera Snap":
                    u_file = st.camera_input(f"Take Photo #{p_idx+1}", key=f"cam_{p_idx}")
                else:
                    u_file = st.file_uploader(
                        f"Upload Photo #{p_idx+1}", 
                        type=["jpg", "jpeg", "png"], 
                        key=f"photo_upload_{p_idx}"
                    )
            
            with p_col2:
                caption = st.text_input(
                    f"Caption / Location #{p_idx+1}", 
                    value=photo_item.get("caption", ""), 
                    placeholder="e.g., Cold storage room - improper raw meat stack height",
                    key=f"photo_cap_{p_idx}"
                )
            
            if u_file is not None:
                uploaded_photos_data.append({"file": u_file, "caption": caption})
                with st.expander(f"Preview Photo #{p_idx+1}", expanded=False):
                    st.image(u_file, width=220)
            
            st.markdown("---")

        st.button("Add Another Photo Evidence Slot", on_click=add_photo_slot)
        st.divider()

        # AUDITOR'S NOTES FIELD
        st.subheader("Auditor's Notes & Field Observations")
        notes_default = st.session_state.get('restored_notes', "")
        auditor_notes = st.text_area(
            "General Auditor Notes, Recommendations, or Store Commendations:",
            value=notes_default,
            placeholder="Enter additional field findings, client requests, positive commendations, or specific auditor suggestions here...",
            height=110,
            key="auditor_notes_field"
        )
        st.divider()

        def get_all_failures():
            current_failures = []
            for module_name, df in edited_modules.items():
                failed_rows = df[df["Fail?"] == True]
                for index, row in failed_rows.iterrows():
                    current_failures.append(f"[{row['Class']}] {row['ID']} {row['Description']} - Notes: {row['Notes']}")
            for finding in st.session_state.custom_findings:
                if finding["note"].strip() != "":
                    lvl_code = finding["level"].split()[0]
                    current_failures.append(f"[{lvl_code}] Custom Finding: {finding['note']}")
            return current_failures

        st.subheader("On-Site Action Guide")
        if st.button("Consult SOPs for Immediate Actions"):
            current_failures = get_all_failures()
            
            if len(current_failures) == 0:
                st.success("No deviations flagged yet! Everything looks good.")
            elif not model:
                st.error("Gemini API Key is missing. Set `GEMINI_API_KEY` in secrets to use AI guidance.")
            else:
                company_standards = read_company_standards(active_folder_slug, current_failures)
                    
                with st.spinner("Consulting SOPs for immediate fixes..."):
                    guide_prompt = f"""
                    You are the FSCO for {establishment_name} - {branch_name}. I am currently auditing the kitchen.
                    Here are the items that just failed: {current_failures}
                    
                    Based STRICTLY on the following company SOPs, tell me what EXACT immediate physical action I need to instruct the staff to take right now to fix these specific issues.
                    Do not give me root causes or long-term preventive actions. Just short, bulleted immediate instructions.
                    
                    SOPs:
                    {company_standards}
                    """
                    try:
                        response = model.generate_content(guide_prompt)
                        st.warning(response.text)
                    except Exception as e:
                        st.error(f"Error consulting SOPs: {e}")

        st.divider()
        st.subheader("Verification & Sign-Off")
        st.write("Sign inside the boundary panel below to authenticate this verification log:")

        with st.container():
            canvas_result = st_canvas(
                fill_color="rgba(255, 255, 255, 0)",  
                stroke_width=3,                      
                stroke_color="#000000",              
                background_color="#FFFFFF",          
                height=150,                          
                width=400,                           
                drawing_mode="freedraw",
                key="fsco_signature_canvas",
                update_streamlit=True
            )

        st.divider()

        if st.button("Process Metrics & Generate Final Report", use_container_width=True):

            deductions = 0
            count_L1, count_L2, count_L3 = 0, 0, 0
            changelog_items = []
            failed_items = get_all_failures()
            failed_items_formatted = "\n".join([f"- {item}" for item in failed_items]) if failed_items else "No violations found."
            
            for item in failed_items:
                if "[L1]" in item: 
                    deductions += 25
                    count_L1 += 1
                elif "[L2]" in item: 
                    deductions += 10
                    count_L2 += 1
                elif "[L3]" in item: 
                    deductions += 2
                    count_L3 += 1
                    
            for finding in st.session_state.custom_findings:
                if finding["note"].strip() != "" and finding["changelog"]:
                    changelog_items.append(finding["note"])
                
            final_score = round((1 - (deductions / BASE_SCORE)) * 100, 2)
            
            if final_score >= 95: rating = "Excellent (95-100%)"
            elif final_score >= 85: rating = "Good (85-94%)"
            elif final_score >= 75: rating = "Okay (75-84%)"
            else: rating = "Needs Improvement (<75%)"
            
            st.session_state.cached_score_data = {
                "deductions": deductions,
                "final_score": final_score,
                "rating": rating,
                "count_L1": count_L1,
                "count_L2": count_L2,
                "count_L3": count_L3
            }
            
            progress_context = ""
            with st.spinner("Analyzing historical trends and saving data..."):
                if gc is None:
                    st.warning("Google Sheets credentials (`gcp_service_account`) are not configured. Skipping database save.")
                    progress_context = "Historical data unavailable (Database connection unconfigured)."
                else:
                    try:
                        sheet = gc.open(active_sheet_name).sheet1
                        existing_data = sheet.get_all_values()
                        
                        previous_score = None
                        for row in reversed(existing_data):
                            if len(row) >= 5:
                                # Clean 6-column schema: [Date, Brand, Branch, Score, Deductions, Violations]
                                if len(row) >= 6 and row[1] == establishment_name and row[2] == branch_name:
                                    try:
                                        previous_score = float(row[3].replace('%', ''))
                                        break
                                    except ValueError:
                                        continue
                                elif len(row) == 5 and row[1] == establishment_name:
                                    try:
                                        previous_score = float(row[2].replace('%', ''))
                                        break
                                    except ValueError:
                                        continue
                                    
                        if previous_score is not None:
                            diff = final_score - previous_score
                            if diff > 0:
                                trend = f"Improved by {diff:.2f}%"
                            elif diff < 0:
                                trend = f"Declined by {abs(diff):.2f}%"
                            else:
                                trend = "Unchanged"
                            progress_context = f"Previous Audit Score ({branch_name}): {previous_score:.2f}% | Current Score: {final_score:.2f}% | Trajectory: {trend}"
                        else:
                            progress_context = f"No previous historical data found for {branch_name}. This is the baseline audit."
                            
                        violations_text = " | ".join(failed_items) if failed_items else "No violations found."
                        sheet.append_row([
                            str(audit_date), 
                            establishment_name, 
                            branch_name, 
                            f"{final_score:.2f}%", 
                            deductions, 
                            violations_text
                        ])
                        st.success(f"Audit row saved securely to sheet: {active_sheet_name}")
                        
                    except Exception as e:
                        st.error(f"Could not save to Google Sheets or fetch history. Diagnostic Error: {e}")
                        progress_context = "Historical data unavailable due to database error."

            changelog_prompt = "\n".join([f"- {item}" for item in changelog_items]) if len(changelog_items) > 0 else "No dynamic updates required for this cycle."
            notes_prompt = auditor_notes.strip() if auditor_notes.strip() else "No additional auditor notes recorded for this cycle."
            
            company_standards = read_company_standards(active_folder_slug, failed_items)
            
            if not model:
                st.error("Gemini API Key missing. Set `GEMINI_API_KEY` in secrets to generate the report.")
            else:
                with st.spinner("Gemini is generating the final official report..."):
                    prompt = f"""
                    You are an expert Lead Food Safety Compliance Officer (FSCO) for {establishment_name} - {branch_name}. 
                    I just finished an audit on {audit_date}. 
                    The final score is {final_score:.2f}% ({deductions} points in deductions). 
                    The specific violations found were: 
                    {failed_items_formatted}
                    
                    CRITICAL RULEBOOK ({establishment_name} PRPs & SOPs):
                    You MUST base your Root Cause analysis and Preventive Actions EXACTLY on these company standards. 
                    Do not invent generic solutions if a solution exists in these rules. 
                    ---
                    {company_standards}
                    ---
                    
                    Write a detailed, highly clinical Audit Executive Summary Report.
                    
                    CRITICAL FORMATTING RULES:
                    - ONLY use **bold** text for section titles or headers.
                    - DO NOT use inline bolding inside of paragraphs.
                    - Use simple dashes (-) instead of em-dashes (—).
                    
                    Format the report STRICTLY with these sections exactly as named:

                    **1. Executive Summary**
                    Objective: Summarize the surveillance audit purpose.
                    Current Compliance Status: State the score.
                    Administrative Breakdown: Hypothesize the root cause of the systemic failures based on our standards.
                    Key Verdict: Give a strict directive for immediate next steps.

                    **2. FSMS Administration: Changelog**
                    Based on the auditor's explicit instructions, here are the items that MUST be added to the manual's changelog:
                    {changelog_prompt}
                    Format this cleanly. If it says "No dynamic updates required", output exactly that.

                    **3. Audit Scoring & Finding Summary**
                    Write exactly: "See quantitative table below."

                    **3.2 Detailed Finding & Resolution Plan**
                    Provide a clean list of the specific violations found.

                    **4. Corrective and Preventive Action (CAPA) Summary**
                    For every L1 and L2 violation, provide a recommended action plan. Number each violation sequentially (e.g., 1., 2., 3.). 
                    Format EXACTLY like this:
                    1. Issue: (State the violation)
                    Immediate Correction: (What to do today)
                    Root Cause: (Hypothesize why it happened)
                    Preventive Action: (How to stop it happening again)

                    **5. Mandatory Compliance Toolkit**
                    List any physical safety equipment that must be procured based on the specific violations.

                    **6. Historical Progress Summary**
                    Data: [{progress_context}]
                    Provide a portfolio-level 1-2 sentence commentary on the facility's compliance trajectory based on the provided data.

                    **7. Auditor Notes & Recommendations**
                    Direct observations and suggestions provided by the Lead Auditor:
                    {notes_prompt}
                    """
                    
                    try:
                        response = model.generate_content(prompt)
                        st.session_state.cached_report_text = response.text.replace('**', '')
                        st.session_state.report_generated = True
                        clear_audit_draft()
                    except Exception as e:
                        st.error(f"Diagnostic Error: {e}")

        if st.session_state.report_generated:
            st.divider()
            
            scores = st.session_state.cached_score_data
            m_col1, m_col2, m_col3 = st.columns(3)
            m_col1.metric("Final Verification Score", f"{scores['final_score']:.2f}%")
            m_col2.metric("Total Point Deductions", f"-{scores['deductions']} pts")
            m_col3.metric("Operational Status", scores['rating'].split()[0])
            
            if scores['final_score'] >= 95:
                st.success("Facility is operating within an optimal compliance parameter.")
            elif scores['final_score'] >= 85:
                st.info("Facility displays stable parameters with minor correction paths needed.")
            else:
                st.error("Escalated Alert: System parameters drop below baseline safety levels.")
            
            st.subheader("Generated Executive Summary")
            st.write(st.session_state.cached_report_text)

            signature_saved = False
            if canvas_result.image_data is not None:
                img_matrix = canvas_result.image_data
                
                if np.any(img_matrix[:, :, 3] > 0): 
                    raw_sketch = Image.fromarray(img_matrix.astype('uint8'), 'RGBA')
                    
                    rgb_signature = Image.new("RGB", raw_sketch.size, (255, 255, 255))
                    rgb_signature.paste(raw_sketch, mask=raw_sketch.split()[3])
                    
                    temp_sig_path = "fsco_signature_temp.png"
                    rgb_signature.save(temp_sig_path, "PNG")
                    signature_saved = True
            
            temp_image_files = [] 
            
            try:
                pdf = FPDF()
                pdf.set_left_margin(10)
                pdf.set_right_margin(10)
                pdf.add_page()
                
                logo_path = None
                for p_logo in ["logo.png", "logo.jpg", "templates/logo.png", "templates/logo.jpg", "assets/logo.png"]:
                    if os.path.exists(p_logo):
                        logo_path = p_logo
                        break

                if logo_path:
                    pdf.image(logo_path, x=150, y=8, w=45)

                # REPORT HEADER
                pdf.set_font("Times", 'B', 14)
                pdf.cell(135, 8, txt="FSCO Monthly Surveillance & Verification Report", ln=True, align='L')
                pdf.set_draw_color(180, 180, 180)
                pdf.line(10, pdf.get_y() + 2, 200, pdf.get_y() + 2)
                pdf.ln(5)
                
                pdf.set_font("Times", 'B', 10)
                pdf.cell(42, 5, "Establishment Name:", ln=False)
                pdf.set_font("Times", '', 10)
                pdf.cell(0, 5, f"{establishment_name}", ln=True)

                pdf.set_font("Times", 'B', 10)
                pdf.cell(42, 5, "Branch / Location:", ln=False)
                pdf.set_font("Times", '', 10)
                pdf.cell(0, 5, f"{branch_name}", ln=True)
                
                pdf.set_font("Times", 'B', 10)
                pdf.cell(42, 5, "Lead Auditor / FSCO:", ln=False)
                pdf.set_font("Times", '', 10)
                pdf.cell(0, 5, f"{fsco_name}", ln=True)
                
                pdf.set_font("Times", 'B', 10)
                pdf.cell(42, 5, "Audit Operation Date:", ln=False)
                pdf.set_font("Times", '', 10)
                pdf.cell(0, 5, f"{audit_date}", ln=True)
                pdf.ln(5)

                # VISUAL SCORE BADGE / BANNER
                score_val = scores['final_score']
                if score_val >= 95.0:
                    bg_r, bg_g, bg_b = 220, 245, 230
                    border_r, border_g, border_b = 40, 167, 69
                    txt_r, txt_g, txt_b = 20, 100, 35
                    status_label = "EXCELLENT (OPTIMAL COMPLIANCE)"
                elif score_val >= 85.0:
                    bg_r, bg_g, bg_b = 255, 243, 205
                    border_r, border_g, border_b = 255, 193, 7
                    txt_r, txt_g, txt_b = 133, 100, 4
                    status_label = "GOOD (STABLE COMPLIANCE)"
                else:
                    bg_r, bg_g, bg_b = 248, 215, 218
                    border_r, border_g, border_b = 220, 53, 69
                    txt_r, txt_g, txt_b = 114, 28, 36
                    status_label = "NEEDS IMPROVEMENT (CRITICAL ACTION REQUIRED)"

                banner_y = pdf.get_y()
                pdf.set_fill_color(bg_r, bg_g, bg_b)
                pdf.set_draw_color(border_r, border_g, border_b)
                pdf.rect(10, banner_y, 190, 10, 'DF')

                pdf.set_xy(10, banner_y + 2)
                pdf.set_font("Times", 'B', 10)
                pdf.set_text_color(txt_r, txt_g, txt_b)
                pdf.cell(190, 6, txt=f"VERIFICATION SCORE: {score_val:.2f}%  |  STATUS: {status_label}", align='C', ln=True)
                pdf.set_text_color(0, 0, 0)
                pdf.ln(4)
                
                safe_text = clean_unicode_text(st.session_state.cached_report_text)
                
                MAIN_SECTION_TITLES = [
                    "1. Executive Summary",
                    "2. FSMS Administration: Changelog",
                    "3. Audit Scoring & Finding Summary",
                    "3.2 Detailed Finding & Resolution Plan",
                    "4. Corrective and Preventive Action (CAPA) Summary",
                    "5. Mandatory Compliance Toolkit",
                    "6. Historical Progress Summary",
                    "7. Auditor Notes & Recommendations"
                ]

                for line in safe_text.split('\n'):
                    line = line.strip()
                    
                    if not line:
                        pdf.ln(2)
                        continue
                    
                    if line.startswith("3.2") or "Detailed Finding" in line:
                        pdf.add_page()
                    
                    # 1. MAIN SECTION HEADERS (GREY BANNER)
                    is_main_header = any(line.startswith(title) for title in MAIN_SECTION_TITLES) or "Audit Scoring" in line
                    
                    if is_main_header:
                        if pdf.get_y() > 230:
                            pdf.add_page()

                        pdf.ln(4)
                        current_y = pdf.get_y()
                        pdf.set_fill_color(240, 240, 240)
                        pdf.rect(10, current_y, 190, 8, 'F')
                        
                        pdf.set_font("Times", 'B', 11)
                        pdf.set_text_color(0, 0, 0)
                        pdf.cell(0, 8, txt=line, ln=True, fill=True)
                        pdf.ln(3)
                        
                        if "3. Audit Scoring" in line or "Finding Summary" in line:
                            pdf.set_font("Times", 'B', 10)
                            pdf.cell(0, 6, txt="3.1 Quantitative Audit Metric Breakdown", ln=True)
                            pdf.set_font("Times", '', 9)
                            pdf.cell(0, 5, txt="Defects weighted across a baseline threshold of 1000 Operational System points.", ln=True)
                            pdf.ln(2)
                            
                            pdf.set_fill_color(220, 225, 230)
                            pdf.set_font("Times", 'B', 9)
                            pdf.cell(60, 7, "Risk Metric Classification", border=1, align='L', fill=True)
                            pdf.cell(40, 7, "Deduction Weight", border=1, align='C', fill=True)
                            pdf.cell(40, 7, "Observed Deviations", border=1, align='C', fill=True)
                            pdf.cell(50, 7, "Total Point Deficit", border=1, align='C', fill=True, ln=True)
                            
                            pdf.set_font("Times", '', 9)
                            pdf.cell(60, 6, "Critical Deviations", border=1)
                            pdf.cell(40, 6, "-25 pts / item", border=1, align='C')
                            pdf.cell(40, 6, str(scores['count_L1']), border=1, align='C')
                            pdf.cell(50, 6, f"-{scores['count_L1'] * 25} pts", border=1, align='C', ln=True)
                            
                            pdf.cell(60, 6, "Major Deviations", border=1)
                            pdf.cell(40, 6, "-10 pts / item", border=1, align='C')
                            pdf.cell(40, 6, str(scores['count_L2']), border=1, align='C')
                            pdf.cell(50, 6, f"-{scores['count_L2'] * 10} pts", border=1, align='C', ln=True)
                            
                            pdf.cell(60, 6, "Minor Deviations", border=1)
                            pdf.cell(40, 6, "-2 pts / item", border=1, align='C')
                            pdf.cell(40, 6, str(scores['count_L3']), border=1, align='C')
                            pdf.cell(50, 6, f"-{scores['count_L3'] * 2} pts", border=1, align='C', ln=True)
                            
                            pdf.set_fill_color(245, 245, 245)
                            pdf.set_font("Times", 'B', 9)
                            pdf.cell(60, 7, "Final Metrics Aggregation", border=1, fill=True)
                            pdf.cell(40, 7, f"Score: {scores['final_score']:.2f}%", border=1, align='C', fill=True)
                            
                            total_observed_count = scores['count_L1'] + scores['count_L2'] + scores['count_L3']
                            pdf.cell(40, 7, f"{total_observed_count} Items", border=1, align='C', fill=True)
                            pdf.cell(50, 7, f"-{scores['deductions']} pts", border=1, align='C', fill=True, ln=True)
                            pdf.ln(4)
                    
                    # 2. CAPA ISSUE HEADERS
                    elif any(line.startswith(f"{i}. Issue:") for i in range(1, 100)) or line.startswith("Issue "):
                        if pdf.get_y() > 240:
                            pdf.add_page()
                        pdf.ln(4)
                        pdf.set_draw_color(200, 200, 200)
                        pdf.line(10, pdf.get_y(), 200, pdf.get_y())
                        pdf.ln(3)
                        
                        parts = line.split(":", 1)
                        pdf.set_font("Times", 'B', 10)
                        pdf.set_text_color(160, 30, 30)
                        pdf.cell(0, 5, txt=parts[0] + ":", ln=True)
                        
                        if len(parts) > 1 and parts[1].strip():
                            pdf.set_font("Times", 'B', 10)
                            pdf.set_text_color(0, 0, 0)
                            pdf.multi_cell(0, 5, txt=parts[1].strip())
                        pdf.ln(2)

                    # 3. SUB-ITEMS WITH LABELS
                    elif ":" in line and len(line.split(":")[0]) < 35:
                        parts = line.split(":", 1)
                        label = parts[0].strip() + ":"
                        body_val = parts[1].strip() if len(parts) > 1 else ""
                        
                        pdf.set_font("Times", 'B', 10)
                        pdf.set_text_color(40, 40, 40)
                        pdf.cell(0, 5, txt=label, ln=True)
                        
                        if body_val:
                            pdf.set_font("Times", '', 10)
                            pdf.set_text_color(0, 0, 0)
                            pdf.multi_cell(0, 5, txt=body_val)
                        pdf.ln(2.5)
                        
                    # 4. STANDARD PARAGRAPHS
                    else:
                        pdf.set_font("Times", '', 10)
                        pdf.set_text_color(0, 0, 0)
                        pdf.multi_cell(0, 5, txt=line)
                        pdf.ln(2)

                # --- RENDER PHOTO EVIDENCE ANNEX IN PDF (2-COLUMN GRID) ---
                if uploaded_photos_data:
                    if pdf.get_y() > 190:
                        pdf.add_page()
                    
                    pdf.ln(4)
                    current_y = pdf.get_y()
                    pdf.set_fill_color(240, 240, 240)
                    pdf.rect(10, current_y, 190, 8, 'F')
                    pdf.set_font("Times", 'B', 11)
                    pdf.set_text_color(0, 0, 0)
                    pdf.cell(0, 8, txt="8. Photographic Evidence Log", ln=True, fill=True)
                    pdf.ln(4)

                    col_x_positions = [12, 105]
                    col_width = 83

                    for p_pair_idx in range(0, len(uploaded_photos_data), 2):
                        pair = uploaded_photos_data[p_pair_idx:p_pair_idx+2]
                        
                        if pdf.get_y() > 180:
                            pdf.add_page()

                        row_start_y = pdf.get_y()
                        max_row_bottom_y = row_start_y

                        for c_offset, photo_item in enumerate(pair):
                            u_file = photo_item["file"]
                            cap_text = photo_item["caption"].strip() if photo_item["caption"].strip() else f"Field Evidence #{p_pair_idx + c_offset + 1}"
                            x_pos = col_x_positions[c_offset]

                            with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as tmp:
                                img = Image.open(u_file)
                                img.convert("RGB").save(tmp.name)
                                tmp_path = tmp.name
                                temp_image_files.append(tmp_path)

                            pdf.set_xy(x_pos, row_start_y)
                            pdf.set_font("Times", 'B', 9)
                            pdf.set_text_color(0, 0, 0)
                            pdf.multi_cell(col_width, 4, txt=clean_unicode_text(f"Exhibit 8.{p_pair_idx + c_offset + 1}: {cap_text}"))
                            
                            caption_end_y = pdf.get_y() + 1

                            with Image.open(tmp_path) as pil_img:
                                img_w, img_h = pil_img.size
                                calc_h = col_width * (img_h / img_w)

                            pdf.image(tmp_path, x=x_pos, y=caption_end_y, w=col_width)
                            col_bottom_y = caption_end_y + calc_h + 4

                            if col_bottom_y > max_row_bottom_y:
                                max_row_bottom_y = col_bottom_y

                        pdf.set_y(max_row_bottom_y)

                # PREVENT ORPHANED SIGNATURE BLOCK
                if signature_saved and os.path.exists("fsco_signature_temp.png"):
                    if pdf.get_y() > 210:
                        pdf.add_page()
                    
                    pdf.ln(6)
                    pdf.set_font("Times", 'B', 10)
                    pdf.set_text_color(0, 0, 0)
                    pdf.cell(0, 5, txt="Authorized Verification Signature:", ln=True)
                    pdf.ln(2)
                    pdf.image("fsco_signature_temp.png", x=12, w=55)
                    pdf.ln(2)
                    pdf.set_font("Times", '', 9)
                    pdf.cell(0, 4, txt=f"{fsco_name} | Certified Food Safety Officer", ln=True)
                    pdf.cell(0, 4, txt="Knife & Ember Food Consultancy Services", ln=True)
                    os.remove("fsco_signature_temp.png") 
                
                # FILENAME FORMAT: Brand (Branch) - Executive Summary Report - [Date].pdf
                safe_est_filename = "".join(c for c in establishment_name if c.isalnum() or c in (' ', '_', '-')).strip()
                safe_branch_filename = "".join(c for c in branch_name if c.isalnum() or c in (' ', '_', '-')).strip()
                
                if safe_branch_filename:
                    pdf_filename = f"{safe_est_filename} ({safe_branch_filename}) - Executive Summary Report - {audit_date}.pdf"
                else:
                    pdf_filename = f"{safe_est_filename} - Executive Summary Report - {audit_date}.pdf"

                pdf.output(pdf_filename)

                for tmp_img in temp_image_files:
                    if os.path.exists(tmp_img):
                        os.remove(tmp_img)
                
                st.markdown("### Server-Side Transmission Dispatch")
                
                with st.spinner(f"Pushing official PDF report directly to {PRIMARY_RECIPIENT_EMAIL}..."):
                    email_success = auto_email_report(
                        recipient_email=PRIMARY_RECIPIENT_EMAIL,
                        pdf_path=pdf_filename,
                        client_name=establishment_name,
                        branch_name=branch_name,
                        score=scores['final_score'],
                        status=scores['rating'].split()[0]
                    )
                    if email_success:
                        st.success(f"Official PDF delivered safely to {PRIMARY_RECIPIENT_EMAIL}")
                
                if os.path.exists(pdf_filename):
                    os.remove(pdf_filename)
                    st.toast("Temporary compilation file cleared from server cache safely.")
                    
            except Exception as e:
                st.error(f"Automated Processing Delivery Error: {e}")

    with tab2:
        st.subheader("Historical Metric Tracker")
        st.write(f"Review compliance trajectory histories synchronized with **{active_sheet_name}**.")
        
        if st.button("Sync Database Records", use_container_width=True):
            if gc is None:
                st.error("Google Sheets connection is unconfigured. Make sure `gcp_service_account` is in your secrets.")
            else:
                with st.spinner(f"Fetching data from '{active_sheet_name}'..."):
                    try:
                        sheet = gc.open(active_sheet_name).sheet1
                        data = sheet.get_all_values()
                        
                        if len(data) > 1: 
                            raw_df = pd.DataFrame(data[1:])
                            num_cols = raw_df.shape[1]
                            
                            # Parse 6-column schema: [Date, Brand, Branch, Score, Deductions, Violations]
                            if num_cols >= 6:
                                df = raw_df.iloc[:, :6]
                                df.columns = ["Date", "Brand", "Branch", "Score", "Deductions", "Violations"]
                            else:
                                df = raw_df.iloc[:, :5]
                                df.columns = ["Date", "Brand", "Score", "Deductions", "Violations"]
                                df["Branch"] = "Main Branch"
                            
                            df["Score"] = df["Score"].astype(str).str.replace("%", "").str.strip()
                            df["Score"] = pd.to_numeric(df["Score"], errors="coerce")
                            df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
                            df = df.dropna(subset=["Score", "Date"]).sort_values(by="Date")
                            
                            unique_branches = ["All Branches"] + sorted(list(df["Branch"].astype(str).unique()))
                            selected_branch_view = st.selectbox("Select Branch View:", options=unique_branches)
                            
                            if selected_branch_view != "All Branches":
                                filtered_df = df[df["Branch"] == selected_branch_view]
                            else:
                                filtered_df = df

                            st.line_chart(data=filtered_df, x="Date", y="Score", use_container_width=True)
                            
                            st.write("**Raw Historical Data Vault:**")
                            st.dataframe(filtered_df, use_container_width=True, hide_index=True)
                            
                        else:
                            st.info("No audit data found in this database sheet yet.")
                            
                    except Exception as e:
                        st.error(f"Could not load dashboard. Diagnostic Error: {e}")

    st.sidebar.button("End Session (Log Out)", on_click=lambda: st.session_state.update(logged_in=False, report_generated=False, cached_report_text="") or st.rerun(), use_container_width=True)